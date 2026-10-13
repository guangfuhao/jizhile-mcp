from __future__ import annotations

import csv
import json
from pathlib import Path

URL_NAMES = ("原始视频URL", "video_url", "videoUrl", "download_url", "url", "视频链接")


def seconds(value) -> float | None:
    if value in (None, ""):
        return None
    try:
        if isinstance(value, (int, float)):
            number = float(value)
        else:
            parts = str(value).strip().split(":")
            number = sum(float(v) * 60**i for i, v in enumerate(reversed(parts)))
        return number if 0 < number <= 3600 else None
    except ValueError:
        return None


def load_items(input_path: str, ids: list[str] | None = None) -> list[dict]:
    path = Path(input_path).expanduser().resolve()
    if not path.is_file() or path.stat().st_size > 20 * 1024 * 1024:
        raise ValueError("输入文件不存在或超过20MB。")
    records = []
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            for sheet in workbook:
                headers = None
                for rownum, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    if rownum > 10050:
                        raise ValueError("表格超过10000行处理上限。")
                    if headers is None:
                        if rownum <= 30 and any(v in URL_NAMES for v in row if isinstance(v, str)):
                            headers = [str(v).strip() if v is not None else "" for v in row]
                        continue
                    item = dict(zip(headers, row))
                    item["source_row"] = rownum; item["source_sheet"] = sheet.title
                    records.append(item)
        finally:
            workbook.close()
    elif path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as source:
            records = [dict(row, source_row=i) for i, row in enumerate(csv.DictReader(source), 2)]
    elif path.suffix.lower() == ".json":
        data = json.loads(path.read_text())
        records = data.get("items", []) if isinstance(data, dict) else data
        if not isinstance(records, list):
            raise ValueError("JSON须为数组或包含items数组。")
    else:
        raise ValueError("支持XLSX、CSV、JSON清单。")
    normalized = []; seen = set()
    for position, row in enumerate(records, 1):
        if not isinstance(row, dict):
            raise ValueError("每条输入须为对象。")
        url = next((str(row[k]).strip() for k in URL_NAMES if row.get(k) and str(row[k]).startswith("https://")), None)
        local = row.get("video_path")
        if not url and not local:
            # Blank spreadsheet footer rows are not video records.
            if not any(row.get(k) for k in ("序号", "id", "视频标识", *URL_NAMES)):
                continue
            raise ValueError(f"第{row.get('source_row', position)}行缺少有效视频直链或video_path。")
        item_id = str(row.get("id", row.get("序号", position)))
        if item_id in seen:
            raise ValueError("存在重复条目ID，无法安全断点续跑。")
        seen.add(item_id)
        if ids is not None and item_id not in ids:
            continue
        normalized.append({"id": item_id, "video_url": url, "video_path": str(Path(local).expanduser().resolve()) if local else None,
                           "duration_hint_seconds": seconds(row.get("duration_seconds", row.get("时长"))),
                           "title": str(row.get("title", row.get("视频标题（去话题词）", "")))[:1000],
                           "account": str(row.get("account", row.get("发布账号", "")))[:300],
                           "source_row": row.get("source_row", position), "source_sheet": row.get("source_sheet")})
    if not normalized or len(normalized) > 500:
        raise ValueError("须选择1至500条有效视频。")
    if ids is not None and set(ids) != {i["id"] for i in normalized}:
        raise ValueError("部分指定ID不存在。")
    return normalized
